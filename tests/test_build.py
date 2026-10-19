import unittest, re, json, subprocess
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
DL=ROOT/'site/downloads'

class V8BuildTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.obs=pd.read_csv(DL/'observations.csv')
        cls.reg=pd.read_csv(DL/'source_registry.csv')
        cls.measures=pd.read_csv(DL/'measure_registry.csv')
        cls.concepts=pd.read_csv(DL/'concept_registry.csv')
        cls.sm=pd.read_csv(DL/'source_measure_map.csv')
        cls.cov=pd.read_csv(DL/'source_concept_coverage.csv')
        cls.sync=pd.read_csv(DL/'synchronisation_audit.csv')
        cls.rel=pd.read_csv(DL/'source_relationships.csv')
        cls.eq=pd.read_csv(DL/'observation_equivalence_groups.csv')
        cls.html=(ROOT/'site/index.html').read_text(encoding='utf-8')
        cls.js=(ROOT/'site/assets/app.js').read_text(encoding='utf-8')

    def test_core_counts(self):
        self.assertEqual(len(self.obs),2248)
        self.assertEqual(self.obs.jurisdiction.nunique(),43)
        self.assertEqual(self.obs.source_dataset.nunique(),43)
        self.assertEqual(len(self.reg),50)
        self.assertEqual(self.obs.measure.nunique(),107)
        self.assertEqual(self.obs.concept_id.nunique(),19)

    def test_observation_ids_unique(self):
        self.assertEqual(len(self.obs),self.obs.observation_id.nunique())

    def test_all_observation_sources_registered(self):
        self.assertFalse(set(self.obs.source_dataset.astype(str))-set(self.reg.source_id.astype(str)))

    def test_all_registry_sources_are_visible_in_coverage(self):
        self.assertEqual(set(self.reg.source_id.astype(str)),set(self.cov.source_id.astype(str)))

    def test_loaded_and_refreshable_sources_are_explicit(self):
        loaded=self.reg[pd.to_numeric(self.reg.loaded_observations,errors='coerce').fillna(0)>0]
        waiting=self.reg[pd.to_numeric(self.reg.loaded_observations,errors='coerce').fillna(0)==0]
        self.assertEqual(len(loaded),43)
        self.assertEqual(len(waiting),7)
        self.assertTrue((waiting.state=='refreshable').all())

    def test_all_measures_mapped_once(self):
        self.assertEqual(set(self.obs.measure.astype(str)),set(self.measures.measure.astype(str)))
        self.assertFalse(self.measures.measure.duplicated().any())
        self.assertTrue(self.measures.concept_id.isin(self.concepts.concept_id).all())

    def test_all_source_measure_pairs_mapped(self):
        pairs=set(map(tuple,self.obs[['source_dataset','measure']].drop_duplicates().astype(str).to_records(index=False)))
        mapped=set(map(tuple,self.sm[['source_dataset','measure']].drop_duplicates().astype(str).to_records(index=False)))
        self.assertEqual(pairs,mapped)

    def test_provenance_complete_for_every_observation(self):
        for col in ['source_url','source_locator','source_record_id','transformation']:
            self.assertTrue(self.obs[col].fillna('').astype(str).str.strip().ne('').all(),col)
        self.assertTrue(self.obs.source_url.astype(str).str.startswith('http').all())
        pa=pd.read_csv(DL/'provenance_audit.csv')
        self.assertEqual(len(pa),len(self.obs))
        self.assertTrue(pa.provenance_complete.astype(str).str.lower().isin(['true','1']).all())

    def test_grade_a_is_definition_verified(self):
        a=self.obs[self.obs.comparability_grade=='A']
        self.assertGreater(len(a),500)
        self.assertTrue(a.definition_verified.astype(str).str.lower().isin(['true','1','yes']).all())

    def test_wpb_values_are_not_mislabeled_as_year_end_stock(self):
        wpb=self.obs[self.obs.source_dataset.astype(str).str.contains('WPB')]
        self.assertGreater(len(wpb),50)
        self.assertFalse((wpb.measure=='stock_eoy_total').any())
        self.assertTrue(wpb[wpb.measure.str.contains('total')].measure.eq('population_point_total').all())

    def test_exact_equivalence_groups_are_exact_only(self):
        if len(self.eq):
            self.assertTrue((pd.to_numeric(self.eq.source_count,errors='coerce')>=2).all())
            self.assertFalse(self.eq.measure.fillna('').str.contains(r' \| ').any())
            self.assertTrue(self.eq.agreement_status.isin(['exact','conflict']).all())

    def test_source_relationships_have_distinct_sources(self):
        if len(self.rel):
            self.assertTrue((self.rel.source_a!=self.rel.source_b).all())
            self.assertFalse(self.rel.duplicated().any())

    def test_synchronisation_audit_has_no_errors(self):
        self.assertFalse((self.sync.severity=='error').any(),self.sync.to_dict('records'))
        required={'all observation sources registered','all observations mapped to concepts','all source-measure pairs mapped','complete observation provenance'}
        self.assertTrue(required.issubset(set(self.sync.check)))

    def test_palestine_is_central_and_complete(self):
        p=self.obs[self.obs.jurisdiction=='Mandatory Palestine']
        self.assertEqual(len(p),1135)
        self.assertGreaterEqual(p.measure.nunique(),50)
        self.assertGreaterEqual(p.facility_name.nunique(),20)
        required={'stock_eoy_total','avg_daily_total','admissions_total','stock_eoy_convicted','stock_eoy_nonconvicted','admissions_convicted','admissions_safe_custody','stock_eoy_sentence_band','disciplinary_offences','recidivism_rate_pct','remarks'}
        self.assertTrue(required.issubset(set(p.measure)))

    def test_palestine_snapshot_cell_fidelity(self):
        wb=load_workbook(ROOT/'pipeline/snapshots/palestine_snapshot.xlsx',data_only=False,read_only=True)
        ws=wb['prison date']
        q=self.obs[(self.obs.source_dataset=='PAL_V2') & self.obs.source_record_id.astype(str).str.startswith('prison date!')]
        checked=0;fail=[]
        def norm(v):
            if v is None:return ''
            if isinstance(v,float) and v.is_integer():return str(int(v))
            return str(v).replace('\r\n','\n').strip()
        for _,r in q.iterrows():
            m=re.search(r'prison date!([A-Z]+\d+)',str(r.source_record_id))
            self.assertIsNotNone(m)
            cell=ws[m.group(1)].value
            raw='' if pd.isna(r.raw_value_text) else str(r.raw_value_text)
            if raw.strip():
                checked+=1
                if norm(raw)!=norm(cell):fail.append((r.source_record_id,raw,cell))
            elif str(r.transformation)=='direct_numeric':
                checked+=1
                try:
                    if abs(float(r.value)-float(cell))>1e-9:fail.append((r.source_record_id,r.value,cell))
                except Exception:fail.append((r.source_record_id,r.value,cell))
        self.assertGreater(checked,750)
        self.assertFalse(fail,fail[:10])

    def test_all_ui_views_exist(self):
        for v in ['palestine','compare','atlas','map','data','quality','downloads']:
            self.assertIn(f'id="view-{v}"',self.html)

    def test_all_palestine_chart_containers_exist(self):
        ids=['pFlowStock','pLegalComposition','pAdmissionsMix','pTurnover','pRates','pAverageStockRatio','pAnnualChange','pCentralStock','pCentralAdmissions','pLockupStock','pLockupAdmissions','pDetention','pFacilityRanking','pFacilityCoverage','pLabour','pLabourShare','pLabourValue','pExpenditure','pDiscipline','pPunishments','pRecidivismRate','pRecidivismMix','pSentences','pEscapesExecutions','pMovements','pLunatics','pVisitors','pNotesTimeline']
        for id_ in ids:self.assertIn(f'id="{id_}"',self.html,id_)

    def test_atlas_represents_all_sources_and_measures(self):
        self.assertIn('id="sourceGrid"',self.html)
        self.assertIn('id="sourceDetail"',self.html)
        self.assertIn('id="measureSelect"',self.html)
        self.assertIn('id="sourceNetwork"',self.html)
        self.assertIn('const S = window.CPE_SOURCES',self.js)
        self.assertIn('const MR = window.CPE_MEASURE_REGISTRY',self.js)

    def test_connectivity_functions_exist(self):
        for fn in ['showObservationDrawer','openSource','openMeasure','renderSourceDetail','renderMeasureAtlas','renderNetwork','setDataFilters']:
            self.assertIn(f'function {fn}',self.js)

    def test_multiyear_functions_exist(self):
        for mode in ['ma3','ma5','periodavg','index','change']:
            self.assertIn(f'value="{mode}"',self.html)
        self.assertIn('function transformPoints',self.js)

    def test_local_static_assets(self):
        required=['site/assets/plotly.min.js','site/assets/app.js','site/assets/styles.css','site/data/observations.js','site/data/sources.js','site/data/measure_registry.js','site/data/concept_registry.js','site/data/source_measure_map.js','site/data/source_concept_coverage.js','site/data/world_outline.js']
        for p in required:self.assertTrue((ROOT/p).exists(),p)

    def test_downloads_expose_mapping_layers(self):
        for f in ['observations.csv','source_registry.csv','measure_registry.csv','concept_registry.csv','source_measure_map.csv','source_concept_coverage.csv','source_relationships.csv','observation_equivalence_groups.csv','jurisdiction_registry.csv','synchronisation_audit.csv']:
            self.assertTrue((DL/f).exists(),f)

if __name__=='__main__': unittest.main()
